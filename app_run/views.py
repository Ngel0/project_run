from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework import status
from rest_framework.viewsets import ModelViewSet, ReadOnlyModelViewSet
from rest_framework.views import APIView
from rest_framework.filters import SearchFilter, OrderingFilter
from django_filters.rest_framework import DjangoFilterBackend
from django.conf import settings
from django.contrib.auth import get_user_model
from django.shortcuts import get_object_or_404
from django.db.models import Count, Sum, Min, Max, Q
from geopy.distance import geodesic
from openpyxl import load_workbook
from .models import Run, AthleteInfo, Challenge, Position, CollectibleItem
from .serializers import RunSerializer, UserSerializer, UserDetailSerializer, PositionSerializer, CollectibleItemSerializer
from .paginations import RunPagination, UserPagination, PositionPagination

User = get_user_model()

@api_view(['GET'])
def company_details(request):
    return Response({
            'company_name':settings.COMPANY_NAME,
            'slogan':settings.SLOGAN,
            'contacts':settings.CONTACTS
        })

class RunViewSet(ModelViewSet):
    queryset = Run.objects.all()
    serializer_class = RunSerializer
    filter_backends = [DjangoFilterBackend, OrderingFilter]
    filterset_fields = ['athlete', 'status']
    ordering_fields = ['created_at']
    pagination_class = RunPagination

    def get_queryset(self):
        return Run.objects.select_related('athlete').all()

class UserViewSet(ReadOnlyModelViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    filter_backends = [SearchFilter, OrderingFilter]
    search_fields = ['last_name', 'first_name']
    ordering_fields = ['date_joined']
    pagination_class = UserPagination

    def get_queryset(self):
        # добавляем аннотацию для подсчета завершенных забегов
        qs = qs.annotate(runs_finished=Count('runs', filter=Q(runs__status='finished')))
        
        qs = self.queryset.filter(is_superuser=False) #сразу исключаем админов
        type = self.request.query_params.get('type', None)
        if type == 'coach':
            return qs.filter(is_staff=True)
        if type == 'athlete':
            return qs.filter(is_staff=False)
        return qs

    def get_serializer_class(self):
        if self.action == 'retrieve':
            return UserDetailSerializer
        return super().get_serializer_class()

class StartRunView(APIView):
    def post(self, request, id):
        run = get_object_or_404(Run, id=id)
        if run.status != 'init':
            return Response({'message':'Can not start the run'}, status=status.HTTP_400_BAD_REQUEST)
        run.status = 'in_progress'
        run.save()
        return Response({'message':run.status}, status=status.HTTP_200_OK)

class StopRunView(APIView):
    def post(self, request, id):
        run = get_object_or_404(Run, id=id)
        if run.status != 'in_progress':
            return Response({'message':'Can not stop the run'}, status=status.HTTP_400_BAD_REQUEST)
        run.status = 'finished'

        # calculating run distance
        coordinates = list(Position.objects.filter(run=id).order_by('id').values_list('latitude', 'longitude'))
        total_distance = 0.0
        if len(coordinates) >= 2:
            total_distance = sum(
                geodesic(coordinates[i], coordinates[i + 1]).kilometers
                for i in range(len(coordinates) - 1)
            )
        run.distance = total_distance
        run.save()

        # checking distance to collectible items
        for position in coordinates:
            for item in CollectibleItem.objects.all():
                if item.users.filter(id=run.athlete.id).exists():
                    continue  # skip if the user already has this item
                if item.latitude is not None and item.longitude is not None:
                    distance_to_item = geodesic(position, (item.latitude, item.longitude)).meters
                    if distance_to_item <= 100:  # within 100 meters
                        item.users.add(run.athlete)

        # calculating run time
        aggrs = Position.objects.filter(run=id).aggregate(max_time=Max('date_time'), min_time=Min('date_time'))
        max_time = aggrs['max_time']
        min_time = aggrs['min_time']
        if max_time and min_time:
            timedelta = max_time - min_time
            run.run_time_seconds = int(timedelta.total_seconds())
            run.save()

        # challenge 10 runs finished
        if Run.objects.filter(athlete=run.athlete, status='finished').count() >= 10:
            Challenge.objects.create(full_name='Сделай 10 Забегов!', athlete=run.athlete)

        # challenge 50 km
        if Run.objects.filter(athlete=run.athlete).aggregate(Sum('distance'))['distance__sum'] >= 50:
            Challenge.objects.create(full_name='Пробеги 50 километров!', athlete=run.athlete)
        
        return Response({'message':run.status}, status=status.HTTP_200_OK)

class AthleteInfoView(APIView):
    def get(self, request, user_id):
        athlete = get_object_or_404(User, id=user_id)
        athlete_info, created = AthleteInfo.objects.get_or_create(user=athlete)
        data = {
            'goals': athlete_info.goals,
            'weight': athlete_info.weight,
            'user_id': user_id,
        }
        return Response(data, status=status.HTTP_200_OK)

    def put(self, request, user_id):
        athlete = get_object_or_404(User, id=user_id)
        weight = request.data.get('weight')
        if weight is not None:
            try:
                weight = int(weight)
            except (ValueError):
                return Response({'message': 'Invalid weight value'}, status=status.HTTP_400_BAD_REQUEST)
            if weight <= 0 or weight >= 900:
                return Response({'message': 'Invalid weight value'}, status=status.HTTP_400_BAD_REQUEST)
        athlete_info, created = AthleteInfo.objects.update_or_create(user=athlete, defaults={
            'goals': request.data.get('goals'),
            'weight': weight
        })
        data = {
            'goals': athlete_info.goals,
            'weight': athlete_info.weight,
            'user_id': user_id,
        }
        return Response(data, status=status.HTTP_201_CREATED)

class ChallengesView(APIView):
    def get(self, request):
        athlete_id = request.GET.get('athlete')
        if athlete_id is None:
            queryset = Challenge.objects.all()
        else:
            try:
                athlete_id = int(athlete_id)
            except(ValueError):
                return Response({'message': 'Invalid athlete id'}, status=status.HTTP_400_BAD_REQUEST)
            queryset = list(Challenge.objects.filter(athlete=athlete_id))
        data = []
        for challenge in queryset:
            item = {'full_name': challenge.full_name, 'athlete': challenge.athlete.id}
            data.append(item)
        return Response(data, status=status.HTTP_200_OK)

class PositionViewSet(ModelViewSet):
    queryset = Position.objects.all()
    serializer_class = PositionSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['run']
    pagination_class = PositionPagination

class CollectibleItemView(APIView):
    def get(self, request):
        items = CollectibleItem.objects.all()
        serializer = CollectibleItemSerializer(items, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

class CollectibleFromFileView(APIView):
    def post(self, request):
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({'message': 'No Excel file was uploaded'}, status=status.HTTP_400_BAD_REQUEST)
        workbook = load_workbook(uploaded_file)
        worksheet = workbook.active
        invalid_rows = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            name, uid, value, latitude, longitude, url = row
            serializer = CollectibleItemSerializer(data={
                'name': name,
                'uid': uid,
                'latitude': latitude,
                'longitude': longitude,
                'picture': url,
                'value': value
            })
            if serializer.is_valid():
                serializer.save()
            else:
                invalid_rows.append(list(row))
        return Response(invalid_rows, status=status.HTTP_200_OK)
